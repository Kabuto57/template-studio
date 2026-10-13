"""Checks that .pptx files are sound: zip integrity, well-formed XML, consistent package
relationships, and (when available) that python-pptx and LibreOffice can open them.
usage: check_pptx.py <file.pptx>...     Set REQUIRE_PPTX_TOOLS=1 to fail when the tools are missing (CI)."""
import os
import posixpath
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile

REQUIRE = os.environ.get("REQUIRE_PPTX_TOOLS") == "1"
try:
    from pptx import Presentation
except ImportError:
    Presentation = None
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
if REQUIRE and (Presentation is None or not SOFFICE):
    sys.exit("REQUIRE_PPTX_TOOLS=1 but python-pptx or LibreOffice is missing")

CT = "{http://schemas.openxmlformats.org/package/2006/content-types}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def check(path):
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None, "corrupt zip member"
        names = set(z.namelist())
        assert len(names) == len(z.namelist()), "duplicate zip entries"
        for n in names:
            if n.endswith((".xml", ".rels")):
                data = z.read(n)
                assert data.count(b"<?xml") <= 1, "%s has two XML declarations" % n
                ET.fromstring(data)
        # every relationship target exists
        for n in names:
            if not n.endswith(".rels"):
                continue
            base = posixpath.dirname(posixpath.dirname(n))
            for rel in ET.fromstring(z.read(n)):
                if rel.get("TargetMode") == "External":
                    continue
                target = rel.get("Target")
                full = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                assert full in names, "%s points to missing part %s" % (n, full)
        # every slide is registered: content type, presentation relationship, sldId >= 256
        overrides = {o.get("PartName").lstrip("/") for o in ET.fromstring(z.read("[Content_Types].xml")) if o.tag == CT + "Override"}
        slides = sorted(n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
        pres_rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(z.read("ppt/_rels/presentation.xml.rels"))}
        ids = ET.fromstring(z.read("ppt/presentation.xml")).find(P + "sldIdLst")
        listed = {posixpath.normpath(posixpath.join("ppt", pres_rels[s.get(R + "id")])) for s in ids}
        assert all(int(s.get("id")) >= 256 for s in ids), "sldId below 256"
        assert len({s.get("id") for s in ids}) == len(list(ids)), "duplicate sldId"
        for s in slides:
            assert s in overrides, "%s has no content-type override" % s
            assert s in listed, "%s is not in the slide list" % s
        assert listed == set(slides), "slide list and slide parts differ"
        for o in overrides:
            assert o in names, "content type for missing part %s" % o
    n_slides = len(slides)
    if Presentation is not None:
        prs = Presentation(path)
        assert len(prs.slides) == n_slides
        for slide in prs.slides:
            for shape in slide.shapes:
                _ = shape.shape_type, shape.has_text_frame and shape.text_frame.text
    if SOFFICE:
        with tempfile.TemporaryDirectory() as tmp:
            # its own profile folder, so checks running in parallel do not fight over one
            subprocess.run([SOFFICE, "-env:UserInstallation=file://%s/profile" % tmp, "--headless", "--convert-to", "pdf", "--outdir", tmp, path],
                           check=True, capture_output=True, timeout=180)
            pdf = os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + ".pdf")
            assert os.path.getsize(pdf) > 1000, "LibreOffice produced no PDF"
    return n_slides


failed = 0
for f in sys.argv[1:]:
    try:
        print("ok   %s (%d slides)" % (os.path.basename(f), check(f)))
    except Exception as exc:  # report every file, then fail
        failed += 1
        print("FAIL %s: %s" % (os.path.basename(f), exc))
print("tools: python-pptx=%s libreoffice=%s" % (Presentation is not None, bool(SOFFICE)))
sys.exit(1 if failed else 0)
